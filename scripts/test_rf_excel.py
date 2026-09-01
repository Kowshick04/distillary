import os
import json
import openpyxl

# Load model json
with open("assets/model/rf_model.json", "r") as f:
    rf_data = json.load(f)

feature_names = rf_data["feature_names"]
print("RF Feature names:", feature_names)

def predict(inputs):
    estimators = rf_data["estimators"]
    classes = rf_data["classes"]
    features = [inputs.get(name, 0.0) for name in feature_names]
    class_probs = [0.0] * len(classes)
    
    for est in estimators:
        left = est["children_left"]
        right = est["children_right"]
        feature = est["feature"]
        threshold = est["threshold"]
        value_dist = est["value"]
        
        curr = 0
        while left[curr] != -1 and right[curr] != -1:
            sf = feature[curr]
            val = features[sf]
            thresh = threshold[curr]
            if val <= thresh:
                curr = left[curr]
            else:
                curr = right[curr]
        
        counts = value_dist[curr][0]
        tot = sum(counts)
        if tot > 0:
            for i, c in enumerate(counts):
                class_probs[i] += (c / tot) / len(estimators)
                
    best_idx = 0
    max_p = -1.0
    for i, p in enumerate(class_probs):
        if p > max_p:
            max_p = p
            best_idx = i
            
    cls_int = classes[best_idx]
    label = 'normal'
    if cls_int == 0: label = 'low_alcohol'
    elif cls_int == 2: label = 'golden'
    
    return label, max_p * 100.0, class_probs

folder = "excel_drops/processed"
for fn in sorted(os.listdir(folder)):
    if fn.endswith(".xlsx") and not "Copy" in fn:
        wb = openpyxl.load_workbook(os.path.join(folder, fn), data_only=True)
        sheet = wb.active
        print(f"\n==================== {fn} ====================")
        
        # Scan rows from bottom up to find 2 fermenters
        seen = {}
        for r in range(sheet.max_row, 6, -1):
            f_num = sheet.cell(row=r, column=5).value
            feed_h = sheet.cell(row=r, column=9).value
            ret_h = sheet.cell(row=r, column=10).value
            setup_sg = sheet.cell(row=r, column=29).value # col 29 vs 30
            final_sg = sheet.cell(row=r, column=30).value
            setup_alc = sheet.cell(row=r, column=32).value
            final_alc = sheet.cell(row=r, column=33).value
            setup_ph = sheet.cell(row=r, column=36).value
            final_ph = sheet.cell(row=r, column=37).value
            
            try:
                f_num_int = int(f_num) if f_num is not None else None
            except:
                f_num_int = None
                
            if f_num_int and f_num_int not in seen and feed_h is not None:
                seen[f_num_int] = {
                    "row": r,
                    "feeding_hours": float(feed_h or 18.5),
                    "retention_hours": float(ret_h or 48.0),
                    "setup_sg": float(setup_sg or 1.085),
                    "final_sg": float(final_sg or 1.010),
                    "sg_diff": abs(float(setup_sg or 1.085) - float(final_sg or 1.010)),
                    "setup_alc": float(setup_alc or 0.0),
                    "final_alc": float(final_alc or 11.5),
                    "setup_ph": float(setup_ph or 4.5),
                    "final_ph": float(final_ph or 4.48),
                }
                if len(seen) >= 2:
                    break
                    
        for fid, vals in seen.items():
            lbl, conf, probs = predict(vals)
            print(f"  Fermenter {fid} (Excel Row {vals['row']}):")
            print(f"    Inputs: FeedH={vals['feeding_hours']}, RetH={vals['retention_hours']}, FinalSG={vals['final_sg']}, FinalAlc={vals['final_alc']}, FinalPH={vals['final_ph']}")
            print(f"    AI Pred: {lbl.upper()} (Confidence: {conf:.1f}%) | Probs: low_alc={probs[0]:.2f}, normal={probs[1]:.2f}, golden={probs[2]:.2f}")
