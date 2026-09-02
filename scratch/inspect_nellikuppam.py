import openpyxl

file_path = r"excel_drops\processed\Nellikuppam_120KLPD_Synthetic_Data.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

print(f"Workbook: {file_path}")
print(f"Sheet Names: {wb.sheetnames}\n")

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"============================================================")
    print(f"SHEET: '{sname}' (Total Rows: {ws.max_row}, Total Cols: {ws.max_column})")
    print(f"============================================================")
    
    if sname == 'Synthetic Data':
        # Print column headers
        print("\n--- Identified Columns in 'Synthetic Data' ---")
        for c in range(1, ws.max_column + 1):
            h_group = ws.cell(4, c).value or ''
            h_unit = ws.cell(5, c).value or ''
            h_name = ws.cell(6, c).value or ''
            if h_name or h_unit or h_group:
                print(f"Col {c:2d}: [{h_group}] {h_name} ({h_unit})")
                
        # Sample row 7 (first data row)
        print("\n--- Sample First Data Row (Row 7) ---")
        for c in range(1, ws.max_column + 1):
            h_name = ws.cell(6, c).value or f"Col_{c}"
            val = ws.cell(7, c).value
            if val is not None and str(val).strip():
                print(f"  • {h_name}: {val}")

    elif sname == 'Calculation sheet':
        print("\n--- Parameters in 'Calculation sheet' ---")
        for r in range(1, ws.max_row + 1):
            row_vals = [str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
            if row_vals:
                print(f"Row {r:2d}: {' | '.join(row_vals)}")
