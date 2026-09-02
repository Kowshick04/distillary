import openpyxl

wb = openpyxl.load_workbook('Bagalkot_60KLPD_Synthetic_Data.xlsx', data_only=True)
ws = wb['Synthetic Data']

print("=== Bagalkot_60KLPD_Synthetic_Data.xlsx - 'Synthetic Data' Columns ===")
for c in range(1, ws.max_column + 1):
    header1 = ws.cell(4, c).value or ''
    header2 = ws.cell(5, c).value or ''
    header3 = ws.cell(6, c).value or ''
    h = f"{header1} | {header2} | {header3}".strip(" |")
    if h:
        print(f"Col {c:2d}: {h}")

print("\n=== 'Calculation sheet' Parameters ===")
ws_calc = wb['Calculation sheet']
for r in range(1, ws_calc.max_row + 1):
    vals = [str(ws_calc.cell(r, c).value) for c in range(1, ws_calc.max_column + 1) if ws_calc.cell(r, c).value is not None]
    if vals:
        print(f"Row {r:2d}: {' | '.join(vals)}")
